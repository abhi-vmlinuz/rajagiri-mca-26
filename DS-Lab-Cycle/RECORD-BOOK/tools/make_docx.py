import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

AUTHOR = "Abhishek Vincent"
ROLL = "03"

BODY_FONT = "Blinker"
CODE_FONT = "Liberation Mono"

PAGE_W = Cm(21.0)
PAGE_H = Cm(29.7)
MARGIN = Inches(1)
TEXT_WIDTH = Pt(451.3)

CODE_SIZE = Pt(9)
CODE_LINE = Pt(11)

CREST = os.path.join("assets", "crest-faded.png")


def load_meta():
    rows = []
    with open("meta.tsv", "r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(line.rstrip("\n").split("\t"))
    return rows


def set_run_font(run, name, size=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)

    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def no_spacing(paragraph):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_rule(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    ppr.append(borders)


def divider(doc):
    paragraph = doc.add_paragraph()
    add_rule(paragraph)
    no_spacing(paragraph)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    body = paragraph._p
    ppr = body.pPr
    index = list(body).index(ppr) + 1 if ppr is not None else 0
    body.insert(index, start)
    body.append(end)


def add_field(paragraph, instruction, cached, size=None, bold=None):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = str(cached)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    for element in (begin, instr, separate, text, end):
        run._r.append(element)

    set_run_font(run, BODY_FONT, size=size, bold=bold)
    return run


def page_borders(section):
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "text")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "7")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    sect_pr.insert_element_before(
        borders,
        "w:lnNumType",
        "w:pgNumType",
        "w:cols",
        "w:formProt",
        "w:vAlign",
        "w:noEndnote",
        "w:titlePg",
        "w:textDirection",
        "w:bidi",
        "w:rtlGutter",
        "w:docGrid",
        "w:printerSettings",
        "w:sectPrChange",
    )


def crest_watermark(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    no_spacing(paragraph)

    run = paragraph.add_run()
    run.add_picture(CREST, width=Pt(270))

    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))

    anchor = OxmlElement("wp:anchor")
    anchor.set("behindDoc", "1")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")
    anchor.set("relativeHeight", "0")

    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)

    for tag in ("positionH", "positionV"):
        position = OxmlElement(f"wp:{tag}")
        position.set("relativeFrom", "page")
        align = OxmlElement("wp:align")
        align.text = "center"
        position.append(align)
        anchor.append(position)

    for tag in ("wp:extent", "wp:effectExtent"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)

    anchor.append(OxmlElement("wp:wrapNone"))

    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)

    drawing.remove(inline)
    drawing.append(anchor)


def page_number_footer(section):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    no_spacing(paragraph)
    add_field(paragraph, "PAGE", "1", size=Pt(10))


def start_numbering(section, start=1):
    sect_pr = section._sectPr
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:start"), str(start))
    sect_pr.insert_element_before(
        pg_num_type,
        "w:cols",
        "w:formProt",
        "w:vAlign",
        "w:noEndnote",
        "w:titlePg",
        "w:textDirection",
        "w:bidi",
        "w:rtlGutter",
        "w:docGrid",
        "w:printerSettings",
        "w:sectPrChange",
    )


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)

    fmt = normal.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def setup_document():
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.27)

    page_borders(section)
    crest_watermark(section)

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.insert_element_before(
        update,
        "w:hdrShapeDefaults",
        "w:footnotePr",
        "w:endnotePr",
        "w:compat",
        "w:docVars",
        "w:rsids",
        "m:mathPr",
        "w:attachedSchema",
        "w:themeFontLang",
        "w:clrSchemeMapping",
        "w:doNotIncludeSubdocsInStats",
        "w:doNotAutoCompressPictures",
        "w:forceUpgrade",
        "w:captions",
        "w:readModeInkLockDown",
        "w:smartTagType",
        "sl:schemaLibrary",
        "w:shapeDefaults",
        "w:doNotEmbedSmartTags",
        "w:decimalSymbol",
        "w:listSeparator",
    )

    return doc


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def cell_border(cell, edge, size="4"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), "000000")
    borders.append(element)


def fix_table_layout(table, widths):
    tbl = table._tbl
    tblpr = tbl.tblPr

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.insert_element_before(
            layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"
        )
    layout.set(qn("w:type"), "fixed")

    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), "0")
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tblpr.insert_element_before(
        margins, "w:tblLook", "w:tblCaption", "w:tblDescription"
    )

    total = sum(int(width.twips) for width in widths)
    width_element = tblpr.find(qn("w:tblW"))
    if width_element is None:
        width_element = OxmlElement("w:tblW")
        tblpr.append(width_element)
    width_element.set(qn("w:type"), "dxa")
    width_element.set(qn("w:w"), str(total))

    for column, width in zip(table.columns, widths):
        column.width = width


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def add_code_lines(doc, path):
    with open(path, "r", encoding="utf-8") as handle:
        lines = handle.read().rstrip("\n").split("\n")

    for line in lines:
        paragraph = doc.add_paragraph()
        no_spacing(paragraph)
        fmt = paragraph.paragraph_format
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = CODE_LINE

        run = paragraph.add_run(line)
        set_run_font(run, CODE_FONT, size=CODE_SIZE)


def add_index(doc, entries, page_map):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    no_spacing(heading)
    run = heading.add_run("INDEX")
    set_run_font(run, BODY_FONT, size=Pt(18), bold=True)

    spacer = doc.add_paragraph()
    no_spacing(spacer)
    spacer.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Inches(0.6), Inches(5.0), Inches(0.6)]
    fix_table_layout(table, widths)

    header_cells = table.rows[0].cells
    for cell, label in zip(header_cells, ("No.", "Program Title", "Page")):
        paragraph = cell.paragraphs[0]
        no_spacing(paragraph)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(11)
        if label == "Page":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run(label)
        set_run_font(run, BODY_FONT, size=Pt(8.5), bold=True)
        cell_border(cell, "bottom", "8")

    repeat_header(table.rows[0])

    for display_number, title, number, date in entries:
        row = table.add_row()
        cells = row.cells

        for cell, width in zip(cells, widths):
            cell.width = width

        first = cells[0].paragraphs[0]
        no_spacing(first)
        first.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        first.paragraph_format.line_spacing = Pt(11)
        run = first.add_run(display_number)
        set_run_font(run, BODY_FONT, size=Pt(8.5))

        second = cells[1].paragraphs[0]
        no_spacing(second)
        second.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        second.paragraph_format.line_spacing = Pt(11)
        run = second.add_run(title)
        set_run_font(run, BODY_FONT, size=Pt(8.5))

        third = cells[2].paragraphs[0]
        no_spacing(third)
        third.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        third.paragraph_format.line_spacing = Pt(11)
        third.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(third, f"PAGEREF prog_{number} \\h", page_map.get(number, ""), size=Pt(8.5))


def add_program(doc, number, date, title, description, page_map, break_page=True):
    display_number = "Project" if int(number) == 57 else str(int(number))

    first = doc.add_paragraph()
    first.paragraph_format.page_break_before = break_page
    no_spacing(first)
    first.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)

    run = first.add_run(f"Program No: {display_number}")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)
    run = first.add_run("\t")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)
    run = first.add_run(f"Date: {date}")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)

    add_bookmark(first, f"prog_{number}", int(number))

    divider(doc)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    run = paragraph.add_run("Program Title:")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)
    run = paragraph.add_run("\u00a0" + title)
    set_run_font(run, BODY_FONT, size=Pt(11))

    divider(doc)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    run = paragraph.add_run("Description:")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(description)
    set_run_font(run, BODY_FONT, size=Pt(11), italic=True)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(210), WD_TAB_ALIGNMENT.LEFT)
    paragraph.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(f"Author: {AUTHOR}")
    set_run_font(run, BODY_FONT, size=Pt(11))
    run = paragraph.add_run(f"\tRoll No: {ROLL}")
    set_run_font(run, BODY_FONT, size=Pt(11))
    run = paragraph.add_run(f"\tDate: {date}")
    set_run_font(run, BODY_FONT, size=Pt(11))

    divider(doc)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("Source Code:")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)

    add_code_lines(doc, os.path.join("programs", f"p{number}.c"))

    divider(doc)

    paragraph = doc.add_paragraph()
    no_spacing(paragraph)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("Output Text:")
    set_run_font(run, BODY_FONT, size=Pt(11), bold=True)

    add_code_lines(doc, os.path.join("outputs", f"p{number}.txt"))


def build(output_path, page_map):
    rows = load_meta()

    doc = setup_document()

    entries = []
    for row in rows:
        number, date, title, description = row
        number = f"{int(number):02d}"
        display_number = "Project" if int(number) == 57 else str(int(number))
        entries.append((display_number, title, number, date))

    add_index(doc, entries, page_map)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    start_numbering(body, 1)
    page_number_footer(body)

    first = True
    for row in rows:
        number, date, title, description = row
        number = f"{int(number):02d}"
        add_program(doc, number, date, title, description, page_map, break_page=not first)
        first = False

    doc.save(output_path)
    return len(rows)


def read_page_map(pdf_path):
    import subprocess

    text = subprocess.run(
        ["pdftotext", "-layout", pdf_path, "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout

    physical = {}
    for index, page in enumerate(text.split("\f"), start=1):
        for line in page.splitlines()[:4]:
            match = re.match(r"\s*Program No:\s*(\d+|Project)\b", line)
            if match:
                key = "57" if match.group(1) == "Project" else f"{int(match.group(1)):02d}"
                physical.setdefault(key, index)
                break

    if not physical:
        return {}

    offset = min(physical.values()) - 1
    return {key: page - offset for key, page in physical.items()}


def main():
    output = sys.argv[1] if len(sys.argv) > 1 else "record.docx"
    page_map_path = "tools/page_map.json"

    page_map = {}
    if os.path.exists(page_map_path):
        with open(page_map_path, "r", encoding="utf-8") as handle:
            page_map = json.load(handle)

    count = build(output, page_map)
    print(f"wrote {output} with {count} entries")


def refresh(output="record.docx"):
    import shutil
    import subprocess

    build(output, {})

    preview = ".docx-preview"
    shutil.rmtree(preview, ignore_errors=True)
    os.makedirs(preview, exist_ok=True)

    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", preview, output],
        check=True,
        capture_output=True,
    )

    page_map = read_page_map(os.path.join(preview, "record.pdf"))
    with open("tools/page_map.json", "w", encoding="utf-8") as handle:
        json.dump(page_map, handle, indent=1)

    count = build(output, page_map)
    print(f"wrote {output} with {count} entries and refreshed page numbers")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--refresh":
        refresh()
    else:
        main()
